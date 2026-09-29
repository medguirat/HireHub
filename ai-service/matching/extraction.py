"""Text extraction from uploaded CVs (PDF, DOCX, plain text)."""

import io
import re
import zipfile

MIN_TEXT_CHARACTERS = 30


class ExtractionError(Exception):
    """`code` is one of: unsupported_format, unreadable, no_text."""

    def __init__(self, code, message):
        super().__init__(message)
        self.code = code


def _detect_format(data, filename):
    # The file signature wins over the extension: renamed files are common.
    if data[:5] == b"%PDF-":
        return "pdf"
    if data[:2] == b"PK":
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if "word/document.xml" in archive.namelist():
                    return "docx"
        except zipfile.BadZipFile:
            pass
    name = (filename or "").lower()
    if name.endswith(".txt"):
        return "txt"
    if name.endswith(".doc"):
        raise ExtractionError("unsupported_format",
                              "Old Word (.doc) files aren't supported. Please save your CV as PDF or DOCX.")
    raise ExtractionError("unsupported_format", "Please upload your CV as a PDF or DOCX file.")


def _pdf_text(data):
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted and not reader.decrypt(""):
            raise ExtractionError("unreadable", "This PDF is password-protected. Please upload an unprotected copy.")
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    except ExtractionError:
        raise
    except (PdfReadError, ValueError, KeyError) as exc:
        raise ExtractionError("unreadable", "This PDF appears to be damaged and couldn't be read.") from exc


def _docx_text(data):
    import docx

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:  # python-docx raises several unrelated types for corrupt files
        raise ExtractionError("unreadable", "This DOCX file appears to be damaged and couldn't be read.") from exc
    lines = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(lines)


def extract_text(data, filename):
    """Returns (text, format). Raises ExtractionError with a user-facing message."""
    fmt = _detect_format(data, filename)
    if fmt == "pdf":
        text = _pdf_text(data)
    elif fmt == "docx":
        text = _docx_text(data)
    else:
        text = data.decode("utf-8", errors="replace")

    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if len(text) < MIN_TEXT_CHARACTERS:
        raise ExtractionError(
            "no_text",
            "We couldn't find any text in this file. If it's a scanned image, please upload a text-based PDF or DOCX.",
        )
    return text, fmt
