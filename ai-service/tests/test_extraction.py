import io

import docx
import pytest

from matching.extraction import ExtractionError, extract_text
from tests.pdf_utils import make_pdf

CV_LINES = ["Amine Trabelsi - Backend Engineer", "Experience: Java, Spring Boot, PostgreSQL", "Jan 2021 - Present"]


def _docx_bytes(paragraphs):
    document = docx.Document()
    for p in paragraphs:
        document.add_paragraph(p)
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Skills"
    table.rows[0].cells[1].text = "Docker, Kubernetes"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_pdf_text_is_extracted():
    text, fmt = extract_text(make_pdf(CV_LINES), "cv.pdf")
    assert fmt == "pdf"
    assert "Spring Boot" in text and "Jan 2021 - Present" in text


def test_docx_paragraphs_and_tables_are_extracted():
    text, fmt = extract_text(_docx_bytes(CV_LINES), "cv.docx")
    assert fmt == "docx"
    assert "PostgreSQL" in text and "Docker, Kubernetes" in text


def test_file_signature_wins_over_extension():
    _, fmt = extract_text(make_pdf(CV_LINES), "cv.docx")
    assert fmt == "pdf"


def test_old_word_files_are_rejected_with_a_clear_message():
    with pytest.raises(ExtractionError) as err:
        extract_text(b"\xd0\xcf\x11\xe0 legacy word", "cv.doc")
    assert err.value.code == "unsupported_format" and "PDF or DOCX" in str(err.value)


def test_images_are_rejected():
    with pytest.raises(ExtractionError) as err:
        extract_text(b"\x89PNG\r\n\x1a\n...", "cv.png")
    assert err.value.code == "unsupported_format"


def test_pdf_without_text_is_reported_as_scanned():
    with pytest.raises(ExtractionError) as err:
        extract_text(make_pdf([]), "scan.pdf")
    assert err.value.code == "no_text"


def test_damaged_pdf_is_reported():
    with pytest.raises(ExtractionError) as err:
        extract_text(b"%PDF-1.4\nthis is not really a pdf", "cv.pdf")
    assert err.value.code in {"unreadable", "no_text"}
